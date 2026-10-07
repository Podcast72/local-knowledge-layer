"""Local, read-only text extraction for document formats."""
from __future__ import annotations

import tempfile
import zipfile
from pathlib import Path

DOC_EXTENSIONS = {'.pdf', '.txt', '.md', '.rtf', '.docx', '.doc', '.xlsx', '.xls', '.odt', '.ods', '.pages', '.numbers'}
LEGACY_EXTENSIONS = {'.doc'}
ZIP_EXTENSIONS = {'.docx', '.xlsx', '.odt', '.ods', '.pages', '.numbers'}


def _safe_zip(path: Path) -> None:
    if path.is_dir():
        return
    if path.suffix.lower() not in ZIP_EXTENSIONS:
        return
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        if len(infos) > 5000 or sum(item.file_size for item in infos) > 100 * 1024 * 1024:
            raise ValueError('Archive too large after decompression')


def _lines(values) -> str | None:
    result = '\n'.join(str(value).strip() for value in values if value is not None and str(value).strip())
    return result or None


def _xlsx(path: Path) -> str | None:
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        lines = [path.stem]
        for sheet in book:
            lines.append(f'Sheet: {sheet.title}')
            for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                cells = [f'{i+1}: {value}' for i, value in enumerate(row) if value is not None and str(value).strip()]
                if cells:
                    lines.append(f'Row {number}: ' + ' | '.join(cells))
        return _lines(lines)
    finally:
        book.close()


def _xls(path: Path) -> str | None:
    import xlrd
    book = xlrd.open_workbook(str(path), on_demand=True)
    try:
        lines = [path.stem]
        for sheet in book.sheets():
            lines.append(f'Sheet: {sheet.name}')
            for number in range(sheet.nrows):
                cells = [f'{i+1}: {value}' for i, value in enumerate(sheet.row_values(number)) if value != '']
                if cells:
                    lines.append(f'Row {number+1}: ' + ' | '.join(cells))
        return _lines(lines)
    finally:
        book.release_resources()


def _odf(path: Path) -> str | None:
    from odf import table, text
    from odf.opendocument import load
    from odf.teletype import extractText
    document = load(str(path))
    lines = [path.stem]
    if path.suffix.lower() == '.odt':
        for element in document.text.getElementsByType(text.H) + document.text.getElementsByType(text.P):
            lines.append(extractText(element))
    else:
        for sheet in document.spreadsheet.getElementsByType(table.Table):
            lines.append(f'Sheet: {sheet.getAttribute("name")}')
            for row_number, row in enumerate(sheet.getElementsByType(table.TableRow), 1):
                cells = []
                for column_number, cell in enumerate(row.getElementsByType(table.TableCell), 1):
                    value = extractText(cell) or cell.getAttribute('value') or cell.getAttribute('formula')
                    if value:
                        cells.append(f'{column_number}: {value}')
                if cells:
                    lines.append(f'Row {row_number}: ' + ' | '.join(cells))
    return _lines(lines)


def _pages(path: Path) -> str | None:
    from pages import Document
    if path.is_dir():
        from .core import ROOT
        temp_dir = ROOT / 'tmp'
        temp_dir.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(dir=temp_dir, suffix='.pages') as stream:
            with zipfile.ZipFile(stream.name, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
                for child in sorted(path.rglob('*')):
                    if child.is_symlink():
                        raise ValueError('Symlink inside Pages package')
                    if child.is_file():
                        archive.write(child, child.relative_to(path))
            return _pages_file(Path(stream.name), path.stem)
    return _pages_file(path, path.stem)


def _pages_file(path: Path, title: str) -> str | None:
    from pages import Document
    document = Document(str(path))
    lines = [title]
    lines.extend(paragraph.text for paragraph in document.paragraphs)
    for number, table in enumerate(document.tables, 1):
        lines.append(f'Table {number}')
        for row in table.rows:
            cells = []
            for cell in row.cells:
                try:
                    value = cell.text
                except NotImplementedError:
                    continue
                if value:
                    cells.append(value)
            if cells:
                lines.append(' | '.join(cells))
    return _lines(lines)


def _numbers(path: Path) -> str | None:
    from numbers_parser import Document
    document = Document(str(path))
    lines = [path.stem]
    for sheet in document.sheets:
        lines.append(f'Sheet: {sheet.name}')
        for table in sheet.tables:
            lines.append(f'Table: {table.name}')
            for number, row in enumerate(table.rows(), 1):
                cells = []
                for column, cell in enumerate(row, 1):
                    value = getattr(cell, 'value', None)
                    formula = getattr(cell, 'formula', None)
                    if value is not None or formula:
                        content = str(value) if value is not None else ''
                        if formula:
                            content += f' formula {formula}'
                        cells.append(f'{column}: {content}')
                if cells:
                    lines.append(f'Row {number}: ' + ' | '.join(cells))
    return _lines(lines)


def extract(path: Path) -> str | None:
    """Extract useful text without opening a source in write mode."""
    suffix = path.suffix.lower()
    if suffix == '.doc':
        return None
    _safe_zip(path)
    if suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        return _lines(f'Page {number}: {text}' for number, page in enumerate(reader.pages, 1) if (text := (page.extract_text() or '').strip()))
    if suffix == '.docx':
        from docx import Document
        document = Document(str(path))
        lines = [paragraph.text for paragraph in document.paragraphs]
        for number, table in enumerate(document.tables, 1):
            lines.append(f'Table {number}')
            lines.extend(' | '.join(cell.text for cell in row.cells if cell.text) for row in table.rows)
        return _lines(lines)
    if suffix == '.xlsx':
        return _xlsx(path)
    if suffix == '.xls':
        return _xls(path)
    if suffix in {'.odt', '.ods'}:
        return _odf(path)
    if suffix == '.rtf':
        from striprtf.striprtf import rtf_to_text
        return rtf_to_text(path.read_text(encoding='utf-8', errors='replace')).strip() or None
    if suffix == '.pages':
        return _pages(path)
    if suffix == '.numbers':
        return _numbers(path)
    raw = path.read_bytes()
    if b'\0' in raw:
        return None
    try:
        decoded = raw.decode('utf-8-sig')
    except UnicodeDecodeError:
        return None
    return decoded.strip() or None
